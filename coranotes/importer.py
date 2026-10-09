"""Importación desde la carpeta de entrada: validación, extracción de texto y deduplicación.

Nunca modifica los archivos originales. No ejecuta macros, no hace OCR, no descarga URLs.
"""
from __future__ import annotations

import io
import json
import logging
import os
import re
import stat
import zipfile
from datetime import date, datetime
from pathlib import Path

from .db import rows
from .records import AppError, Conflict, NotFound, RegistroIn, InvalidData
from .util import clean_text, new_id, normalize_key, sha256_bytes, utc_iso

log = logging.getLogger("coranotes.importacion")

TEXT_EXT = {".txt", ".md", ".json"}
DOC_EXT = {".docx", ".pdf"}
IMG_EXT = {".png", ".jpg", ".jpeg"}
ALL_EXT = TEXT_EXT | DOC_EXT | IMG_EXT
MAX_NAME = 180
MAX_DEPTH = 4
REPARSE = getattr(stat, "FILE_ATTRIBUTE_REPARSE_POINT", 0x400)

KEYMAP = {
    "fecha": "fecha", "fecha de actividad": "fecha", "dia": "fecha",
    "proyecto": "proyecto",
    "tarea": "tarea_titulo", "tarea titulo": "tarea_titulo",
    "tipo avance": "tipo_avance", "tipo de avance": "tipo_avance",
    "estado problema": "estado_problema", "estado del problema": "estado_problema",
    "actividad": "actividad", "actividad realizada": "actividad",
    "avance": "avance", "avance o resultado": "avance", "resultado": "avance", "resultados": "avance",
    "problemas": "problemas", "problema": "problemas",
    "solucion": "solucion", "solucion o estado del problema": "solucion", "estado de la solucion": "solucion",
    "proximos pasos": "proximos_pasos", "siguientes pasos": "proximos_pasos",
}
META_KEYS = {"fecha", "proyecto", "tarea_titulo", "tipo_avance", "estado_problema"}
SECTION_KEYS = {"actividad", "avance", "problemas", "solucion", "proximos_pasos"}
TIPO_MAP = {"exploracion": "exploracion", "en curso": "en_curso", "validado": "validado"}
ESTADO_MAP = {"sin problema": "sin_problema", "ninguno": "sin_problema", "abierto": "abierto",
              "en curso": "en_curso", "resuelto": "resuelto"}
_KV_RE = re.compile(r"^\s*([A-Za-zÁÉÍÓÚÜÑáéíóúüñ_ ]{2,40})\s*:\s*(.*)$")
_HEAD_RE = re.compile(r"^\s*#{1,6}\s+(.+?)\s*:?\s*$")


class ImportProblem(Exception):
    pass


def is_reparse(path: Path) -> bool:
    try:
        st = os.lstat(path)
    except OSError:
        return False
    return stat.S_ISLNK(st.st_mode) or bool(getattr(st, "st_file_attributes", 0) & REPARSE)


def ignored_name(name: str) -> bool:
    low = name.lower()
    return (name.startswith(("~$", ".", "~")) or low.endswith((".tmp", ".part", ".crdownload", ".partial", ".swp"))
            or low in ("desktop.ini", "thumbs.db"))


def validate_inbox_folder(path_str: str, forbidden: list[Path]) -> Path:
    raw = Path(clean_text(path_str, 400)).expanduser()
    if not raw.is_absolute():
        raise AppError("Indicá una ruta completa, por ejemplo C:\\Users\\vos\\Documents\\CoraNotes Entrada")
    if not raw.is_dir():
        raise AppError("La carpeta no existe")
    if is_reparse(raw):
        raise AppError("No se aceptan enlaces simbólicos ni junctions como carpeta de entrada")
    real = raw.resolve(strict=True)
    if real.parent == real:
        raise AppError("No se puede supervisar una unidad completa. Elegí una carpeta exclusiva.")
    home = Path.home().resolve()
    broad = {home, home / "Desktop", home / "Documents", home / "Downloads", home / "OneDrive",
             home / "Escritorio", home / "Documentos", home / "Descargas"}
    if real in broad:
        raise AppError("Elegí una carpeta exclusiva para CoraNotes, no una carpeta general del usuario.")
    for f in forbidden:
        f = f.resolve()
        if real == f or real.is_relative_to(f) or f.is_relative_to(real):
            raise AppError("La carpeta de entrada no puede contener ni estar dentro de las carpetas de la app "
                           "(código, datos, historial o adjuntos).")
    return real


def parse_date(value) -> date | None:
    if not value:
        return None
    s = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def parse_structured_text(text: str) -> dict:
    lines = text.split("\n")
    fields: dict[str, str] = {}
    i = 0
    if lines and lines[0].strip() == "---":
        for j in range(1, min(len(lines), 60)):
            if lines[j].strip() == "---":
                for line in lines[1:j]:
                    m = _KV_RE.match(line)
                    if m and KEYMAP.get(normalize_key(m.group(1))) in META_KEYS:
                        fields[KEYMAP[normalize_key(m.group(1))]] = m.group(2).strip().strip("\"'")
                i = j + 1
                break
    else:
        while i < len(lines) and i < 20:
            line = lines[i].strip()
            if not line:
                i += 1
                if fields:
                    break
                continue
            m = _KV_RE.match(line)
            key = KEYMAP.get(normalize_key(m.group(1))) if m else None
            if key in META_KEYS:
                fields[key] = m.group(2).strip()
                i += 1
            else:
                break
    sections: dict[str, list[str]] = {}
    current = "_intro"
    for line in lines[i:]:
        h = _HEAD_RE.match(line)
        if h and KEYMAP.get(normalize_key(h.group(1))) in SECTION_KEYS:
            current = KEYMAP[normalize_key(h.group(1))]
            continue
        kv = _KV_RE.match(line)
        if kv and KEYMAP.get(normalize_key(kv.group(1))) in SECTION_KEYS:
            current = KEYMAP[normalize_key(kv.group(1))]
            if kv.group(2).strip():
                sections.setdefault(current, []).append(kv.group(2).strip())
            continue
        if kv and KEYMAP.get(normalize_key(kv.group(1))) in META_KEYS and KEYMAP[normalize_key(kv.group(1))] not in fields:
            fields[KEYMAP[normalize_key(kv.group(1))]] = kv.group(2).strip()
            continue
        sections.setdefault(current, []).append(line)
    for key, buf in sections.items():
        value = clean_text("\n".join(buf))
        if not value:
            continue
        if key == "_intro":
            fields["actividad"] = (value + ("\n" + fields["actividad"] if fields.get("actividad") else ""))
        else:
            fields[key] = (fields.get(key, "") + "\n" + value).strip()
    return fields


def parse_json(data: bytes) -> dict:
    try:
        obj = json.loads(data.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ImportProblem("JSON inválido o no está en UTF-8") from exc
    if not isinstance(obj, dict):
        raise ImportProblem("El JSON debe ser un objeto con los campos de la plantilla")
    out = {}
    for k, v in obj.items():
        key = KEYMAP.get(normalize_key(str(k)))
        if not key:
            continue
        if v is None:
            continue
        if not isinstance(v, (str, int, float)):
            raise ImportProblem(f"El campo «{k}» debe ser texto")
        out[key] = str(v)
    return out


def decode_text(data: bytes) -> str:
    if b"\x00" in data:
        raise ImportProblem("El archivo parece binario (contiene bytes nulos)")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            return data.decode("cp1252")
        except UnicodeDecodeError as exc:
            raise ImportProblem("Codificación no soportada: guardá el archivo como UTF-8") from exc


def extract_docx(data: bytes, limits) -> str:
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise ImportProblem("No es un DOCX válido") from exc
    with zf:
        infos = zf.infolist()
        if len(infos) > 1000:
            raise ImportProblem("El DOCX tiene demasiadas partes internas")
        if sum(i.file_size for i in infos) > limits.docx_bytes:
            raise ImportProblem("El DOCX descomprimido supera el límite configurado")
        for i in infos:
            if i.compress_size and i.file_size / max(i.compress_size, 1) > 200:
                raise ImportProblem("Proporción de compresión sospechosa en el DOCX")
        names = {n.lower() for n in zf.namelist()}
        if "word/document.xml" not in names:
            raise ImportProblem("No es un DOCX válido")
        if any(n.endswith("vbaproject.bin") for n in names):
            raise ImportProblem("El documento contiene macros: no se importa")
    try:
        import docx  # python-docx (parser XML sin resolución de entidades)
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ImportProblem("No se pudo leer el DOCX") from exc
    parts: list[str] = []
    total = 0
    for para in document.paragraphs:
        style = (para.style.name if para.style is not None else "") or ""
        text = para.text
        if style.lower().startswith(("heading", "título", "titulo")) and text.strip():
            text = "## " + text
        parts.append(text)
        total += len(text)
        if total > limits.max_caracteres_extraidos:
            break
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)[: limits.max_caracteres_extraidos]


def extract_pdf(data: bytes, limits) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ImportProblem("El PDF está protegido con contraseña")
        if len(reader.pages) > limits.pdf_paginas:
            raise ImportProblem(f"El PDF supera el límite de {limits.pdf_paginas} páginas")
        parts, total = [], 0
        for page in reader.pages:
            text = page.extract_text() or ""
            parts.append(text)
            total += len(text)
            if total > limits.max_caracteres_extraidos:
                break
    except ImportProblem:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ImportProblem("No se pudo leer el PDF") from exc
    text = "\n".join(parts)[: limits.max_caracteres_extraidos]
    if not text.strip():
        raise ImportProblem("El PDF no tiene texto seleccionable (no se hace OCR)")
    return text


def normalize_fields(fields: dict) -> tuple[dict, list[str]]:
    notes = []
    out = {k: clean_text(v, 10_000) for k, v in fields.items()}
    if "tipo_avance" in out:
        mapped = TIPO_MAP.get(normalize_key(out["tipo_avance"]))
        if not mapped:
            notes.append("tipo de avance desconocido (se usó «en curso»)")
        out["tipo_avance"] = mapped or "en_curso"
    if "estado_problema" in out:
        mapped = ESTADO_MAP.get(normalize_key(out["estado_problema"]))
        if not mapped:
            notes.append("estado del problema desconocido (se usó «sin problema»)")
        out["estado_problema"] = mapped or "sin_problema"
    out["proyecto"] = clean_text(out.get("proyecto", ""), 120)
    out["tarea_titulo"] = clean_text(out.get("tarea_titulo", ""), 200)
    return out, notes


class Importer:
    def __init__(self, db, clock, settings, records, evidence, notices, forbidden_dirs: list[Path]):
        self.db = db
        self.clock = clock
        self.settings = settings
        self.records = records
        self.evidence = evidence
        self.notices = notices
        self.forbidden = forbidden_dirs

    def inbox(self) -> Path | None:
        cfg = self.settings.get("entrada")
        if not cfg.activa or not cfg.carpeta:
            return None
        try:
            return validate_inbox_folder(cfg.carpeta, self.forbidden)
        except AppError:
            return None

    # ---------- registro de documentos ----------
    def _upsert_doc(self, c, rel: str, nombre: str, ext: str, **fields) -> dict:
        now = utc_iso(self.clock.now())
        row = c.execute("SELECT * FROM documentos WHERE ruta_relativa = ?", (rel,)).fetchone()
        if row is None:
            did = new_id("DOC")
            c.execute("INSERT INTO documentos(id, ruta_relativa, nombre, extension, estado, creado_utc, actualizado_utc) "
                      "VALUES (?,?,?,?,?,?,?)", (did, rel, nombre, ext, fields.get("estado", "error"), now, now))
            row = c.execute("SELECT * FROM documentos WHERE id = ?", (did,)).fetchone()
        if fields:
            fields["actualizado_utc"] = now
            fields["origen_presente"] = 1
            cols = ", ".join(f"{k} = ?" for k in fields)
            c.execute(f"UPDATE documentos SET {cols} WHERE id = ?", (*fields.values(), row["id"]))
        return dict(c.execute("SELECT * FROM documentos WHERE id = ?", (row["id"],)).fetchone())

    def _new_version(self, c, doc: dict, digest: str, size: int) -> int:
        version = (doc.get("version") or 0) + 1
        c.execute("INSERT INTO documento_versiones(documento_id, version, sha256, tamano, detectado_utc) "
                  "VALUES (?,?,?,?,?)", (doc["id"], version, digest, size, utc_iso(self.clock.now())))
        return version

    def record_error(self, rel: str, motivo: str) -> None:
        nombre = Path(rel).name
        with self.db.tx() as c:
            self._upsert_doc(c, rel, nombre[:MAX_NAME], Path(rel).suffix.lower(), estado="error", motivo=motivo)

    # ---------- proceso principal ----------
    def process_file(self, path: Path) -> dict:
        inbox = self.inbox()
        if inbox is None:
            return {"estado": "inactiva"}
        path = Path(os.path.abspath(path))
        try:
            rel_path = path.relative_to(inbox)
        except ValueError:
            log.warning("Se ignoró un archivo fuera de la carpeta de entrada")
            return {"estado": "fuera_de_carpeta"}
        rel = rel_path.as_posix()
        name = path.name
        if any(ignored_name(p) for p in rel_path.parts):
            return {"estado": "ignorado"}
        ext = path.suffix.lower()
        try:
            data = self._read_checked(path, inbox, rel_path, ext)
        except ImportProblem as exc:
            self.record_error(rel, str(exc))
            return {"estado": "error", "motivo": str(exc)}
        except FileNotFoundError:
            return {"estado": "desaparecido"}
        digest = sha256_bytes(data)
        limits = self.settings.get("limites")
        with self.db.tx() as c:
            existing = c.execute("SELECT * FROM documentos WHERE ruta_relativa = ?", (rel,)).fetchone()
            existing = dict(existing) if existing else None
            if existing and existing["sha256"] == digest and existing["estado"] != "error":
                if not existing["origen_presente"]:
                    c.execute("UPDATE documentos SET origen_presente = 1 WHERE id = ?", (existing["id"],))
                return {"estado": "sin_cambios", "documento_id": existing["id"]}
            if existing and existing["sha256"] == digest and existing["estado"] == "error":
                return {"estado": "error", "documento_id": existing["id"], "motivo": existing["motivo"]}
            dup = c.execute("SELECT id FROM documentos WHERE sha256 = ? AND ruta_relativa <> ? AND estado IN "
                            "('importado','pendiente_clasificar','evidencia','conflicto') LIMIT 1",
                            (digest, rel)).fetchone()
            if dup and not (existing and existing.get("registro_id")):
                doc = self._upsert_doc(c, rel, name, ext, estado="duplicado", sha256=digest, tamano=len(data),
                                       duplicado_de=dup["id"], motivo=f"Contenido idéntico a {dup['id']}")
                return {"estado": "duplicado", "documento_id": doc["id"]}
            if ext in IMG_EXT:
                doc = self._upsert_doc(c, rel, name, ext, estado="evidencia", sha256=digest, tamano=len(data), motivo="")
                version = self._new_version(c, doc, digest, len(data))
                try:
                    ev = self.evidence.add(name, data, registro_id=None, origen="carpeta_entrada",
                                           documento_id=doc["id"], conn=c)
                except AppError as exc:
                    c.execute("UPDATE documentos SET estado='error', motivo=? WHERE id=?", (exc.message, doc["id"]))
                    return {"estado": "error", "documento_id": doc["id"], "motivo": exc.message}
                c.execute("UPDATE documentos SET evidencia_id=?, version=? WHERE id=?", (ev["id"], version, doc["id"]))
                return {"estado": "evidencia", "documento_id": doc["id"], "evidencia_id": ev["id"]}
            try:
                if ext == ".json":
                    fields = parse_json(data)
                    texto = json.dumps(fields, ensure_ascii=False, indent=1)
                else:
                    if ext in (".txt", ".md"):
                        texto = decode_text(data)
                    elif ext == ".docx":
                        texto = extract_docx(data, limits)
                    else:
                        texto = extract_pdf(data, limits)
                    texto = clean_text(texto)[: limits.max_caracteres_extraidos]
                    fields = parse_structured_text(texto)
                if not texto.strip() or not any(clean_text(fields.get(k, "")) for k in SECTION_KEYS):
                    raise ImportProblem("El archivo no tiene contenido de actividad")
            except ImportProblem as exc:
                self._upsert_doc(c, rel, name, ext, estado="error", motivo=str(exc), sha256=digest, tamano=len(data))
                return {"estado": "error", "motivo": str(exc)}
            fields, notes = normalize_fields(fields)
            return self._apply(c, rel, name, ext, digest, len(data), texto, fields, notes, existing)

    def _read_checked(self, path: Path, inbox: Path, rel_path: Path, ext: str) -> bytes:
        if len(path.name) > MAX_NAME or any(ord(ch) < 32 for ch in str(rel_path)):
            raise ImportProblem("Nombre de archivo demasiado largo o con caracteres no permitidos")
        if len(rel_path.parts) > MAX_DEPTH:
            raise ImportProblem("El archivo está en subcarpetas demasiado profundas")
        if ext not in ALL_EXT:
            raise ImportProblem("Formato no soportado (se aceptan .txt, .md, .json, .docx, .pdf, .png, .jpg)")
        cur = inbox
        for part in rel_path.parts:
            cur = cur / part
            if is_reparse(cur):
                raise ImportProblem("Se rechazan enlaces simbólicos y junctions")
        real = path.resolve(strict=True)
        if not real.is_relative_to(inbox):
            raise ImportProblem("La ruta sale de la carpeta autorizada")
        if not real.is_file():
            raise ImportProblem("No es un archivo común")
        limits = self.settings.get("limites")
        limit = limits.texto_bytes if ext in TEXT_EXT else limits.documento_bytes
        size = real.stat().st_size
        if size == 0:
            raise ImportProblem("El archivo está vacío")
        if size > limit:
            raise ImportProblem(f"Supera el límite de tamaño ({limit // 1024} KB)")
        try:
            with open(real, "rb") as fh:
                data = fh.read(limit + 1)
        except PermissionError as exc:
            raise ImportProblem("No se pudo leer el archivo (en uso o sin permiso)") from exc
        if len(data) > limit:
            raise ImportProblem("Supera el límite de tamaño")
        return data

    def _apply(self, c, rel, name, ext, digest, size, texto, fields, notes, existing) -> dict:
        fecha = parse_date(fields.get("fecha"))
        missing = []
        if not fecha:
            missing.append("fecha válida (AAAA-MM-DD o DD/MM/AAAA)")
        if not fields.get("proyecto"):
            missing.append("proyecto")
        if not fields.get("actividad"):
            missing.append("actividad")
        meta = {k: v for k, v in fields.items()}
        doc = self._upsert_doc(c, rel, name, ext, sha256=digest, tamano=size, texto=texto[:200_000],
                               estado=existing["estado"] if existing else "pendiente_clasificar")
        version = self._new_version(c, doc, digest, size)
        linked = None
        if doc.get("registro_id"):
            linked = c.execute("SELECT * FROM registros WHERE id = ?", (doc["registro_id"],)).fetchone()
        prev_meta = json.loads(doc.get("metadatos") or "{}")
        if missing:
            meta["faltantes"] = missing
            motivo = "Pendiente de clasificar: falta " + ", ".join(missing)
            if linked:
                motivo += f". El registro {linked['id']} conserva la versión anterior."
            meta["revision_importada"] = prev_meta.get("revision_importada")
            c.execute("UPDATE documentos SET estado='pendiente_clasificar', motivo=?, metadatos=?, version=? WHERE id=?",
                      (motivo, json.dumps(meta, ensure_ascii=False), version, doc["id"]))
            return {"estado": "pendiente_clasificar", "documento_id": doc["id"]}
        data = RegistroIn(fecha=fecha, proyecto=fields["proyecto"], tarea_titulo=fields.get("tarea_titulo", ""),
                          actividad=fields.get("actividad", ""), avance=fields.get("avance", ""),
                          problemas=fields.get("problemas", ""), solucion=fields.get("solucion", ""),
                          proximos_pasos=fields.get("proximos_pasos", ""),
                          tipo_avance=fields.get("tipo_avance", "en_curso"),
                          estado_problema=fields.get("estado_problema", "sin_problema"), estado="guardado")
        motivo = "; ".join(notes)
        if linked and not linked["archivado"]:
            if prev_meta.get("revision_importada") and linked["revision"] != prev_meta["revision_importada"]:
                meta["revision_importada"] = prev_meta["revision_importada"]
                c.execute("UPDATE documentos SET estado='conflicto', motivo=?, metadatos=?, version=? WHERE id=?",
                          (f"El documento cambió pero el registro {linked['id']} fue editado en la app: no se "
                           "sobrescribió. Revisá y copiá los cambios a mano.", json.dumps(meta, ensure_ascii=False),
                           version, doc["id"]))
                return {"estado": "conflicto", "documento_id": doc["id"], "registro_id": linked["id"]}
            rec, _ = self.records.update(linked["id"], data, motivo="documento_actualizado", conn=c)
        elif linked and linked["archivado"]:
            c.execute("UPDATE documentos SET estado='importado', motivo=?, version=? WHERE id=?",
                      (f"El documento cambió, pero su registro {linked['id']} está en la papelera.", version, doc["id"]))
            return {"estado": "importado", "documento_id": doc["id"], "registro_id": linked["id"]}
        else:
            rec, _ = self.records.create(data, origen="importado", documento_id=doc["id"], conn=c)
        meta["revision_importada"] = rec["revision"]
        c.execute("UPDATE documentos SET estado='importado', motivo=?, registro_id=?, metadatos=?, version=? WHERE id=?",
                  (motivo, rec["id"], json.dumps(meta, ensure_ascii=False), version, doc["id"]))
        return {"estado": "importado", "documento_id": doc["id"], "registro_id": rec["id"]}

    # ---------- acciones del usuario ----------
    def classify(self, doc_id: str, fecha: date, proyecto: str, tarea_titulo: str = "") -> dict:
        with self.db.tx() as c:
            row = c.execute("SELECT * FROM documentos WHERE id = ?", (doc_id,)).fetchone()
            if not row:
                raise NotFound("No existe ese documento")
            doc = dict(row)
            if doc["estado"] != "pendiente_clasificar":
                raise Conflict("Ese documento no está pendiente de clasificar")
            meta = json.loads(doc["metadatos"] or "{}")
            actividad = meta.get("actividad") or clean_text(doc["texto"], 10_000)
            if not actividad:
                raise InvalidData("El documento no tiene texto de actividad")
            try:
                data = RegistroIn(fecha=fecha, proyecto=clean_text(proyecto, 120),
                                  tarea_titulo=clean_text(tarea_titulo or meta.get("tarea_titulo", ""), 200),
                                  actividad=actividad[:10_000], avance=meta.get("avance", "")[:10_000],
                                  problemas=meta.get("problemas", "")[:10_000], solucion=meta.get("solucion", "")[:10_000],
                                  proximos_pasos=meta.get("proximos_pasos", "")[:10_000],
                                  tipo_avance=meta.get("tipo_avance", "en_curso"),
                                  estado_problema=meta.get("estado_problema", "sin_problema"), estado="guardado")
            except ValueError as exc:
                raise InvalidData(str(exc)) from exc
            linked = None
            if doc.get("registro_id"):
                linked = c.execute("SELECT * FROM registros WHERE id = ? AND archivado = 0",
                                   (doc["registro_id"],)).fetchone()
            if linked:
                rec, _ = self.records.update(linked["id"], data, motivo="documento_clasificado", conn=c)
            else:
                rec, _ = self.records.create(data, origen="importado", documento_id=doc_id, conn=c)
            meta.update({"fecha": fecha.isoformat(), "proyecto": data.proyecto, "revision_importada": rec["revision"]})
            meta.pop("faltantes", None)
            c.execute("UPDATE documentos SET estado='importado', motivo='Clasificado manualmente', registro_id=?, "
                      "metadatos=?, actualizado_utc=? WHERE id=?",
                      (rec["id"], json.dumps(meta, ensure_ascii=False), utc_iso(self.clock.now()), doc_id))
        return {"documento_id": doc_id, "registro_id": rec["id"]}

    def ignore(self, doc_id: str) -> None:
        with self.db.tx() as c:
            cur = c.execute("UPDATE documentos SET estado='ignorado', motivo='Ignorado por el usuario', "
                            "actualizado_utc=? WHERE id=?", (utc_iso(self.clock.now()), doc_id))
            if not cur.rowcount:
                raise NotFound("No existe ese documento")

    def archive_linked_record(self, doc_id: str) -> dict:
        with self.db.read() as c:
            row = c.execute("SELECT registro_id FROM documentos WHERE id = ?", (doc_id,)).fetchone()
        if not row or not row["registro_id"]:
            raise NotFound("Ese documento no tiene un registro asociado")
        return self.records.archive(row["registro_id"], motivo="archivado_origen_eliminado")

    def list(self, estado: str | None = None) -> list[dict]:
        sql = ("SELECT id, ruta_relativa, nombre, extension, estado, motivo, version, registro_id, evidencia_id, "
               "duplicado_de, metadatos, origen_presente, actualizado_utc, substr(texto, 1, 400) AS extracto "
               "FROM documentos")
        params: list = []
        if estado:
            sql += " WHERE estado = ?"
            params.append(estado)
        with self.db.read() as c:
            out = rows(c.execute(sql + " ORDER BY actualizado_utc DESC LIMIT 300", params))
        for d in out:
            d["metadatos"] = json.loads(d["metadatos"] or "{}")
            d["origen_presente"] = bool(d["origen_presente"])
        return out

    # ---------- escaneo completo ----------
    def iter_candidates(self, inbox: Path):
        stack = [(inbox, 0)]
        while stack:
            folder, depth = stack.pop()
            try:
                entries = list(os.scandir(folder))
            except OSError:
                continue
            for e in entries:
                p = Path(e.path)
                if ignored_name(e.name):
                    continue
                if is_reparse(p):
                    rel = p.relative_to(inbox).as_posix()
                    self.record_error(rel, "Se rechazan enlaces simbólicos y junctions")
                    continue
                if e.is_dir(follow_symlinks=False):
                    if depth + 1 < MAX_DEPTH:
                        stack.append((p, depth + 1))
                elif e.is_file(follow_symlinks=False):
                    yield p

    def scan(self, process=None) -> dict:
        inbox = self.inbox()
        if inbox is None:
            return {"estado": "inactiva"}
        seen = set()
        results = []
        for p in self.iter_candidates(inbox):
            seen.add(p.relative_to(inbox).as_posix())
            results.append((process or self.process_file)(p))
        self._detect_removed(seen)
        return {"estado": "ok", "procesados": len(results)}

    def _detect_removed(self, present: set[str]) -> None:
        removed = []
        with self.db.tx() as c:
            for d in rows(c.execute("SELECT id, ruta_relativa, registro_id FROM documentos WHERE origen_presente = 1")):
                if d["ruta_relativa"] not in present:
                    c.execute("UPDATE documentos SET origen_presente = 0, actualizado_utc = ? WHERE id = ?",
                              (utc_iso(self.clock.now()), d["id"]))
                    if d["registro_id"]:
                        removed.append(d)
        for d in removed:
            self.notices.add("origen_eliminado", "Archivo quitado de la carpeta de entrada",
                             f"Se quitó un archivo importado. Su registro {d['registro_id']} sigue guardado; "
                             "si querés, archivalo desde Historial > Documentos.", "#/historial", nativo=False)
